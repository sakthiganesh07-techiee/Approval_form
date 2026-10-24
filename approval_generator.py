from pathlib import Path
from datetime import datetime
from docx import Document

BASE_DIR = Path(__file__).resolve().parent

TEMPLATE_DOCX = BASE_DIR / "approval_template.docx"
OUTPUT_DOCX = BASE_DIR / "approval_form.docx"

CIRCULAR_TEMPLATE_DOCX = BASE_DIR / "circular_template.docx"
CIRCULAR_OUTPUT_DOCX = BASE_DIR / "circular.docx"


def replace_in_paragraph(paragraph, template_values):
    original_text = paragraph.text
    updated_text = original_text

    for placeholder, value in template_values.items():
        updated_text = updated_text.replace(
            placeholder,
            str(value)
        )

    if updated_text != original_text:
        if paragraph.runs:
            paragraph.runs[0].text = updated_text

            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(updated_text)


def replace_in_table(table, template_values):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(
                    paragraph,
                    template_values
                )

            for nested_table in cell.tables:
                replace_in_table(
                    nested_table,
                    template_values
                )


def fill_document(template_path, output_path, template_values, document_name):
    if not template_path.exists():
        raise FileNotFoundError(
            f"{document_name} template not found: {template_path}"
        )

    document = Document(template_path)

    # Replace placeholders in paragraphs.
    for paragraph in document.paragraphs:
        replace_in_paragraph(
            paragraph,
            template_values
        )

    # Replace placeholders in tables, including nested tables.
    for table in document.tables:
        replace_in_table(
            table,
            template_values
        )

    try:
        document.save(output_path)

    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = output_path.with_name(
            f"{output_path.stem}_{timestamp}{output_path.suffix}"
        )
        document.save(output_path)

    print(f"\n{document_name} generated successfully!")
    print(f"Output file: {output_path.resolve()}")

    return output_path


def fill_approval_form(template_values):
    return fill_document(
        TEMPLATE_DOCX,
        OUTPUT_DOCX,
        template_values,
        "Approval form"
    )


def fill_circular(template_values):
    return fill_document(
        CIRCULAR_TEMPLATE_DOCX,
        CIRCULAR_OUTPUT_DOCX,
        template_values,
        "Circular"
    )